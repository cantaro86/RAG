import json
import stat

import pytest
from openpyxl import load_workbook

from evaluation.cli import main as evaluation_main
from evaluation.export_excel import _classification, export_run_to_excel
from evaluation.tracing import derive_trace_diagnostics, write_gzip_json

pytestmark = pytest.mark.evaluation


def _trace(question_id, line_number, question, *, completed=True):
    if not completed:
        return {
            "schema_version": 2,
            "question_id": question_id,
            "question_line": line_number,
            "raw_question": question,
            "status": "failed",
            "error": {"type": "RuntimeError", "message": "graph failed"},
            "documents": {},
            "steps": [
                {
                    "sequence": 0,
                    "checkpoint_step": -1,
                    "nodes": ["__input__"],
                    "next_nodes": ["sanitize_question"],
                    "state": {"question": question, "social_intent": None, "guardrail_status": None},
                    "tasks": [],
                }
            ],
            "diagnostics": {
                "node_path": [],
                "question_versions": [{"sequence": -1, "nodes": ["__input__"], "question": question}],
                "retrieval_attempts": [],
                "reformulation_count": 0,
                "maximum_rewrite_count": 0,
                "terminal_node": None,
                "next_nodes": ["sanitize_question"],
                "failed_nodes": [],
                "social_intent": None,
                "guardrail_status": None,
                "retrieval_succeeded": False,
                "final_has_docs": None,
            },
        }

    document_id = "doc-123"
    sanitized = "Domanda sanitizzata"
    transformed = "Domanda trasformata"
    trace = {
        "schema_version": 2,
        "question_id": question_id,
        "question_line": line_number,
        "raw_question": question,
        "status": "completed",
        "documents": {
            document_id: {
                "page_content": "Contesto finale",
                "metadata": {"source": "source.md", "rerank_score": 0.91},
            }
        },
        "steps": [
            {
                "sequence": 0,
                "checkpoint_step": -1,
                "nodes": ["__input__"],
                "next_nodes": ["sanitize_question"],
                "state": {"question": question},
                "tasks": [],
            },
            {
                "sequence": 1,
                "checkpoint_step": 0,
                "nodes": ["sanitize_question"],
                "next_nodes": ["social_intent"],
                "state": {"question": sanitized, "social_intent": None, "guardrail_status": None},
                "tasks": [],
            },
            {
                "sequence": 2,
                "checkpoint_step": 1,
                "nodes": ["social_intent"],
                "next_nodes": ["init_first_question"],
                "state": {"question": sanitized, "social_intent": "DOMANDA", "guardrail_status": None},
                "tasks": [],
            },
            {
                "sequence": 3,
                "checkpoint_step": 2,
                "nodes": ["init_first_question"],
                "next_nodes": ["domain_guardrail"],
                "state": {
                    "question": sanitized,
                    "social_intent": "DOMANDA",
                    "guardrail_status": None,
                    "first_question": True,
                },
                "tasks": [],
            },
            {
                "sequence": 4,
                "checkpoint_step": 3,
                "nodes": ["domain_guardrail"],
                "next_nodes": ["retrieve_and_filter"],
                "state": {
                    "question": sanitized,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                },
                "tasks": [],
            },
            {
                "sequence": 5,
                "checkpoint_step": 4,
                "nodes": ["retrieve_and_filter"],
                "next_nodes": ["transform_query"],
                "state": {
                    "question": sanitized,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                    "documents": [],
                    "has_docs": False,
                    "rewrite_count": 1,
                },
                "tasks": [],
            },
            {
                "sequence": 6,
                "checkpoint_step": 5,
                "nodes": ["transform_query"],
                "next_nodes": ["retrieve_and_filter"],
                "state": {
                    "question": transformed,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                    "documents": [],
                    "has_docs": False,
                    "rewrite_count": 1,
                },
                "tasks": [],
            },
            {
                "sequence": 7,
                "checkpoint_step": 6,
                "nodes": ["retrieve_and_filter"],
                "next_nodes": ["generate_with_docs"],
                "state": {
                    "question": transformed,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                    "documents": [{"$document": document_id}],
                    "has_docs": True,
                    "rewrite_count": 0,
                },
                "tasks": [],
            },
            {
                "sequence": 8,
                "checkpoint_step": 7,
                "nodes": ["generate_with_docs"],
                "next_nodes": ["clean_answer"],
                "state": {
                    "question": transformed,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                    "documents": [{"$document": document_id}],
                    "has_docs": True,
                    "rewrite_count": 0,
                    "generation": "Risposta finale",
                },
                "tasks": [],
            },
            {
                "sequence": 9,
                "checkpoint_step": 8,
                "nodes": ["clean_answer"],
                "next_nodes": [],
                "state": {
                    "question": transformed,
                    "social_intent": "DOMANDA",
                    "guardrail_status": "ON_TOPIC",
                    "documents": [{"$document": document_id}],
                    "has_docs": True,
                    "rewrite_count": 0,
                    "generation": "Risposta finale",
                },
                "tasks": [],
            },
        ],
    }
    trace["diagnostics"] = derive_trace_diagnostics(question, trace["steps"])
    return trace


def _non_rag_trace(question_id, line_number, question):
    steps = [
        {
            "sequence": 0,
            "checkpoint_step": -1,
            "nodes": ["__input__"],
            "next_nodes": ["sanitize_question"],
            "state": {"question": question},
            "tasks": [],
        },
        {
            "sequence": 1,
            "checkpoint_step": 0,
            "nodes": ["sanitize_question"],
            "next_nodes": ["social_intent"],
            "state": {
                "question": question,
                "documents": [],
                "social_intent": None,
                "guardrail_status": None,
            },
            "tasks": [],
        },
        {
            "sequence": 2,
            "checkpoint_step": 1,
            "nodes": ["social_intent"],
            "next_nodes": ["handle_hello"],
            "state": {
                "question": question,
                "documents": [],
                "social_intent": "SALUTO",
                "guardrail_status": None,
            },
            "tasks": [],
        },
        {
            "sequence": 3,
            "checkpoint_step": 2,
            "nodes": ["handle_hello"],
            "next_nodes": [],
            "state": {
                "question": question,
                "documents": [],
                "social_intent": "SALUTO",
                "guardrail_status": None,
                "generation": "Ciao, sono un agente IA.",
            },
            "tasks": [],
        },
    ]
    return {
        "schema_version": 2,
        "question_id": question_id,
        "question_line": line_number,
        "raw_question": question,
        "status": "completed",
        "documents": {},
        "steps": steps,
        "diagnostics": derive_trace_diagnostics(question, steps),
    }


def _write_run(run_dir):
    traces_dir = run_dir / "traces"
    traces_dir.mkdir(parents=True)
    completed_id = "q0001-abcd1234"
    non_rag_id = "q0002-efgh5678"
    failed_id = "q0003-ijkl9012"
    write_gzip_json(traces_dir / f"{completed_id}.json.gz", _trace(completed_id, 1, "Domanda originale"))
    write_gzip_json(traces_dir / f"{non_rag_id}.json.gz", _non_rag_trace(non_rag_id, 2, "Ciao"))
    write_gzip_json(
        traces_dir / f"{failed_id}.json.gz",
        _trace(failed_id, 3, "Domanda fallita", completed=False),
    )
    samples = [
        {
            "schema_version": 2,
            "id": completed_id,
            "response_type": "rag",
            "user_input": "Domanda originale",
            "response": "Risposta finale",
            "retrieved_contexts": ["Contesto finale"],
            "trace_path": f"traces/{completed_id}.json.gz",
        },
        {
            "schema_version": 2,
            "id": non_rag_id,
            "response_type": "non_rag",
            "user_input": "Ciao",
            "response": "Ciao, sono un agente IA.",
            "retrieved_contexts": [],
            "trace_path": f"traces/{non_rag_id}.json.gz",
        },
    ]
    scores = [
        {
            "schema_version": 2,
            "id": completed_id,
            "response_type": "rag",
            "metrics": {"faithfulness": 0.8, "context_utilization": 0.7, "answer_relevancy": 0.9},
            "errors": {},
            "skipped": {},
        },
        {
            "schema_version": 2,
            "id": non_rag_id,
            "response_type": "non_rag",
            "metrics": {"faithfulness": None, "context_utilization": None, "answer_relevancy": 0.6},
            "errors": {},
            "skipped": {
                "faithfulness": "No retrieved contexts",
                "context_utilization": "No retrieved contexts",
            },
        },
    ]
    (run_dir / "samples.jsonl").write_text("".join(json.dumps(sample) + "\n" for sample in samples), encoding="utf-8")
    (run_dir / "scores.jsonl").write_text("".join(json.dumps(score) + "\n" for score in scores), encoding="utf-8")
    (run_dir / "run.json").write_text(
        json.dumps({"schema_version": 2, "run_id": "test-run", "status": "scored", "question_count": 3}) + "\n",
        encoding="utf-8",
    )
    (run_dir / "summary.json").write_text(
        json.dumps({"schema_version": 2, "collection": {"questions": 3, "completed": 2, "failed": 1}}) + "\n",
        encoding="utf-8",
    )


def test_excel_export_combines_artifacts_and_keeps_failed_questions(tmp_path):
    run_dir = tmp_path / "run"
    _write_run(run_dir)

    output_path = export_run_to_excel(run_dir)

    assert output_path == run_dir / "evaluation_results.xlsx"
    assert stat.S_IMODE(output_path.stat().st_mode) == 0o600
    workbook = load_workbook(output_path, read_only=False, data_only=True)
    try:
        assert workbook.sheetnames == ["Results", "Run Summary", "Legend"]
        worksheet = workbook["Results"]
        headers = [cell.value for cell in worksheet[1]]
        completed = dict(zip(headers, (cell.value for cell in worksheet[2]), strict=True))
        non_rag = dict(zip(headers, (cell.value for cell in worksheet[3]), strict=True))
        failed = dict(zip(headers, (cell.value for cell in worksheet[4]), strict=True))

        assert worksheet.freeze_panes == "A2"
        assert worksheet.tables["EvaluationResults"] is not None
        assert len(headers) == 21
        assert {
            "question_id",
            "scoring_status",
            "classification_raw",
            "question_versions",
            "contexts_truncated",
            "scoring_errors",
            "reformulation_count",
            "maximum_rewrite_count",
            "terminal_node",
            "collection_error",
            "trace_path",
        }.isdisjoint(headers)
        assert completed["question_line"] == 1
        assert completed["response_type"] == "rag"
        assert completed["classification"] == "on_topic"
        assert completed["social_intent"] == "DOMANDA"
        assert completed["guardrail_status"] == "ON_TOPIC"
        assert completed["sanitized_question"] == "Domanda sanitizzata"
        assert completed["transformed_question"] == "Domanda trasformata"
        assert completed["final_response"] == "Risposta finale"
        assert completed["retrieved_contexts"] == "[Context 1]\nContesto finale"
        assert completed["retrieved_sources"] == "[1] source.md"
        assert completed["faithfulness"] == 0.8
        assert completed["context_utilization"] == 0.7
        assert completed["answer_relevancy"] == 0.9
        assert non_rag["question_line"] == 2
        assert non_rag["response_type"] == "non_rag"
        assert non_rag["classification"] == "greeting"
        assert non_rag["social_intent"] == "SALUTO"
        assert non_rag["guardrail_status"] is None
        assert non_rag["final_response"] == "Ciao, sono un agente IA."
        assert non_rag["retrieved_context_count"] == 0
        assert non_rag["answer_relevancy"] == 0.6
        assert failed["question_line"] == 3
        assert failed["collection_status"] == "failed"
        assert failed["response_type"] is None
        assert failed["classification"] is None
        assert failed["social_intent"] is None
        assert failed["guardrail_status"] is None
        assert failed["final_response"] is None
    finally:
        workbook.close()


def test_excel_export_requires_overwrite_for_existing_report(tmp_path):
    run_dir = tmp_path / "run"
    _write_run(run_dir)
    assert evaluation_main(["export-excel", str(run_dir)]) == 0
    output_path = run_dir / "evaluation_results.xlsx"

    with pytest.raises(FileExistsError, match="Pass --overwrite"):
        export_run_to_excel(run_dir)

    assert export_run_to_excel(run_dir, overwrite=True) == output_path


def test_excel_export_requires_schema_version_2(tmp_path):
    run_dir = tmp_path / "run"
    _write_run(run_dir)
    manifest = json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    manifest["schema_version"] = 1
    (run_dir / "run.json").write_text(json.dumps(manifest) + "\n", encoding="utf-8")

    with pytest.raises(ValueError, match="schema version 1; expected 2"):
        export_run_to_excel(run_dir)


def test_excel_export_supports_collected_run_without_scores_or_summary(tmp_path):
    run_dir = tmp_path / "run"
    _write_run(run_dir)
    (run_dir / "scores.jsonl").unlink()
    (run_dir / "summary.json").unlink()

    output_path = export_run_to_excel(run_dir, tmp_path / "collected.xlsx")

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        rows = workbook["Results"].iter_rows(values_only=True)
        headers = next(rows)
        completed = dict(zip(headers, next(rows), strict=True))
        assert completed["faithfulness"] is None
        assert completed["context_utilization"] is None
        assert completed["answer_relevancy"] is None
    finally:
        workbook.close()


@pytest.mark.parametrize(("social_intent", "expected"), [("SALUTO", "greeting"), ("GRAZIE", "thanks")])
def test_excel_classification_uses_terminal_social_intents(social_intent, expected):
    diagnostics = {"guardrail_status": None, "social_intent": social_intent}

    assert _classification({}, diagnostics) == expected
    assert _classification({}, {**diagnostics, "social_intent": "DOMANDA"}) is None
