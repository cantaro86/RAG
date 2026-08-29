from __future__ import annotations

import json
import re
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from evaluation.tracing import (
    RESPONSE_TYPES,
    load_jsonl,
    read_gzip_json,
    require_schema_version,
    validate_sample_trace,
)

EXCEL_CELL_LIMIT = 32_767
EXCEL_TRUNCATION_NOTICE = "\n\n[TRUNCATED: Excel cell limit reached]"
ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")
METRIC_NAMES = ("faithfulness", "context_utilization", "answer_relevancy")
CLASSIFICATION_LABELS = {
    "ON_TOPIC": "on_topic",
    "OFF_TOPIC": "off_topic",
    "SALUTO": "greeting",
    "GRAZIE": "thanks",
}

COLUMNS = (
    ("question_line", "Original nonblank question line number.", 14),
    ("collection_status", "Trace outcome: completed or failed.", 18),
    ("response_type", "Response category used for reporting: rag or non_rag.", 16),
    ("classification", "Derived route class: on_topic, off_topic, greeting, or thanks.", 16),
    ("social_intent", "Final social_intent output: SALUTO, GRAZIE, or DOMANDA.", 16),
    ("guardrail_status", "Final domain_guardrail output: ON_TOPIC or OFF_TOPIC.", 18),
    ("user_input", "Original question submitted to the graph.", 45),
    ("sanitized_question", "Question after sanitize_question.", 45),
    ("transformed_question", "Last question produced by transform_query, if it ran.", 45),
    ("final_question", "Question in the final available graph checkpoint.", 45),
    ("final_response", "Final response returned by the graph.", 65),
    ("retrieval_succeeded", "True when any retrieval attempt accepted documents.", 18),
    ("retrieval_attempt_count", "Number of retrieve_and_filter executions.", 20),
    ("retrieved_context_count", "Number of final contexts supplied to RAGAS.", 20),
    ("retrieved_sources", "Source metadata for final retrieved documents.", 40),
    ("retrieved_rerank_scores", "Reranker scores for final retrieved documents, when present.", 24),
    ("retrieved_contexts", "Final contexts supplied to RAGAS, numbered and separated by blank lines.", 80),
    ("faithfulness", "RAGAS faithfulness score.", 16),
    ("context_utilization", "RAGAS context utilization score.", 20),
    ("answer_relevancy", "RAGAS answer relevancy score.", 18),
    ("scoring_skipped", "Per-metric reasons for skipped scores.", 40),
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
STATUS_FILLS = {
    "completed": PatternFill("solid", fgColor="E2F0D9"),
    "failed": PatternFill("solid", fgColor="FCE4D6"),
}
CLASSIFICATION_FILLS = {
    "on_topic": PatternFill("solid", fgColor="E2F0D9"),
    "off_topic": PatternFill("solid", fgColor="FCE4D6"),
    "greeting": PatternFill("solid", fgColor="DDEBF7"),
    "thanks": PatternFill("solid", fgColor="FFF2CC"),
}


def _load_json_object(path: Path, *, required: bool) -> dict[str, Any]:
    if not path.is_file():
        if required:
            raise FileNotFoundError(f"Evaluation artifact not found: {path}")
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError(f"Expected a JSON object in {path}")
    require_schema_version(value, path)
    return value


def _index_records(records: Iterable[Mapping[str, Any]], source: Path) -> dict[str, Mapping[str, Any]]:
    indexed: dict[str, Mapping[str, Any]] = {}
    for record in records:
        require_schema_version(record, source)
        record_id = record.get("id")
        if not isinstance(record_id, str) or not record_id:
            raise ValueError(f"Record in {source} has no string id")
        if record_id in indexed:
            raise ValueError(f"Duplicate question id in {source}: {record_id}")
        indexed[record_id] = record
    return indexed


def _trace_index(run_dir: Path) -> dict[str, dict[str, Any]]:
    traces_dir = run_dir / "traces"
    if not traces_dir.is_dir():
        raise FileNotFoundError(f"Trace directory not found: {traces_dir}")

    traces: dict[str, dict[str, Any]] = {}
    for trace_path in sorted(traces_dir.glob("*.json.gz")):
        trace = read_gzip_json(trace_path)
        require_schema_version(trace, trace_path)
        question_id = trace.get("question_id")
        if not isinstance(question_id, str) or not question_id:
            raise ValueError(f"Trace has no string question_id: {trace_path}")
        if question_id in traces:
            raise ValueError(f"Duplicate trace question id: {question_id}")
        traces[question_id] = trace
    if not traces:
        raise ValueError(f"No question traces found in {traces_dir}")
    return traces


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _state_after_node(trace: Mapping[str, Any], node: str) -> Mapping[str, Any]:
    result: Mapping[str, Any] = {}
    for step in _list(trace.get("steps")):
        if not isinstance(step, Mapping) or node not in _list(step.get("nodes")):
            continue
        state = step.get("state")
        if isinstance(state, Mapping):
            result = state
    return result


def _final_state(trace: Mapping[str, Any]) -> Mapping[str, Any]:
    for step in reversed(_list(trace.get("steps"))):
        if isinstance(step, Mapping) and isinstance(step.get("state"), Mapping):
            return step["state"]
    return {}


def _question_from_state(state: Mapping[str, Any]) -> str | None:
    question = state.get("question")
    return question if isinstance(question, str) else None


def _classification(trace: Mapping[str, Any], diagnostics: Mapping[str, Any]) -> str | None:
    raw = diagnostics.get("guardrail_status")
    social = diagnostics.get("social_intent")
    if (not isinstance(raw, str) or not raw) and social in {"SALUTO", "GRAZIE"}:
        raw = social
    if not isinstance(raw, str) or not raw:
        for step in reversed(_list(trace.get("steps"))):
            if not isinstance(step, Mapping):
                continue
            state = _mapping(step.get("state"))
            candidate = state.get("guardrail_status")
            if not isinstance(candidate, str) or not candidate:
                social = state.get("social_intent")
                candidate = social if social in {"SALUTO", "GRAZIE"} else None
            if isinstance(candidate, str) and candidate:
                raw = candidate
                break
    if not isinstance(raw, str) or not raw:
        return None
    normalized = raw.strip().upper()
    return CLASSIFICATION_LABELS.get(normalized, normalized.casefold())


def _document_details(trace: Mapping[str, Any]) -> tuple[str, str]:
    documents = _mapping(trace.get("documents"))
    state_documents = _list(_final_state(trace).get("documents"))
    sources: list[str] = []
    rerank_scores: list[str] = []
    for number, reference in enumerate(state_documents, 1):
        if not isinstance(reference, Mapping) or not isinstance(reference.get("$document"), str):
            continue
        payload = _mapping(documents.get(reference["$document"]))
        metadata = _mapping(payload.get("metadata"))
        source = metadata.get("source")
        if isinstance(source, str) and source:
            sources.append(f"[{number}] {source}")
        score = metadata.get("rerank_score")
        if isinstance(score, int | float) and not isinstance(score, bool):
            rerank_scores.append(f"[{number}] {score:.6g}")
    return "\n".join(sources), "\n".join(rerank_scores)


def _format_contexts(contexts: list[str]) -> str:
    value = "\n\n".join(f"[Context {number}]\n{context}" for number, context in enumerate(contexts, 1))
    value = ILLEGAL_EXCEL_CHARACTERS.sub("", value)
    if len(value) <= EXCEL_CELL_LIMIT:
        return value
    return value[: EXCEL_CELL_LIMIT - len(EXCEL_TRUNCATION_NOTICE)] + EXCEL_TRUNCATION_NOTICE


def _excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_EXCEL_CHARACTERS.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        value = f"'{value}"
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[: EXCEL_CELL_LIMIT - len(EXCEL_TRUNCATION_NOTICE)] + EXCEL_TRUNCATION_NOTICE
    return value


def _format_issue_mapping(value: Any) -> str:
    if not isinstance(value, Mapping):
        return ""
    lines = []
    for name, detail in value.items():
        if isinstance(detail, Mapping):
            issue_type = detail.get("type")
            message = detail.get("message")
            text = ": ".join(str(part) for part in (issue_type, message) if part)
        else:
            text = str(detail)
        lines.append(f"{name}: {text}")
    return "\n".join(lines)


def _result_row(
    trace: Mapping[str, Any],
    sample: Mapping[str, Any] | None,
    score: Mapping[str, Any] | None,
) -> dict[str, Any]:
    diagnostics = _mapping(trace.get("diagnostics"))
    final_state = _final_state(trace)
    sanitized_question = _question_from_state(_state_after_node(trace, "sanitize_question"))
    transformed_question = _question_from_state(_state_after_node(trace, "transform_query"))
    classification = _classification(trace, diagnostics)
    sample = sample or {}
    score = score or {}
    metrics = _mapping(score.get("metrics"))
    contexts_value = sample.get("retrieved_contexts")
    contexts = [context for context in _list(contexts_value) if isinstance(context, str)]
    formatted_contexts = _format_contexts(contexts)
    sources, rerank_scores = _document_details(trace)
    raw_question = trace.get("raw_question")
    user_input = sample.get("user_input")
    if not isinstance(user_input, str):
        user_input = raw_question if isinstance(raw_question, str) else None

    return {
        "question_line": trace.get("question_line"),
        "collection_status": trace.get("status"),
        "response_type": sample.get("response_type"),
        "classification": classification,
        "social_intent": diagnostics.get("social_intent"),
        "guardrail_status": diagnostics.get("guardrail_status"),
        "user_input": user_input,
        "sanitized_question": sanitized_question,
        "transformed_question": transformed_question,
        "final_question": _question_from_state(final_state),
        "final_response": sample.get("response"),
        "retrieval_succeeded": diagnostics.get("retrieval_succeeded"),
        "retrieval_attempt_count": len(_list(diagnostics.get("retrieval_attempts"))),
        "retrieved_context_count": len(contexts),
        "retrieved_sources": sources,
        "retrieved_rerank_scores": rerank_scores,
        "retrieved_contexts": formatted_contexts,
        **{metric: metrics.get(metric) for metric in METRIC_NAMES},
        "scoring_skipped": _format_issue_mapping(score.get("skipped")),
    }


def build_result_rows(run_dir: Path) -> list[dict[str, Any]]:
    run_dir = run_dir.expanduser().resolve()
    samples_path = run_dir / "samples.jsonl"
    samples = _index_records(load_jsonl(samples_path), samples_path)
    scores_path = run_dir / "scores.jsonl"
    scores = _index_records(load_jsonl(scores_path), scores_path) if scores_path.is_file() else {}
    traces = _trace_index(run_dir)

    orphan_samples = sorted(set(samples) - set(traces))
    if orphan_samples:
        raise ValueError(f"Samples have no corresponding trace: {', '.join(orphan_samples)}")
    for question_id, sample in samples.items():
        response_type = sample.get("response_type")
        if response_type not in RESPONSE_TYPES:
            raise ValueError(f"Sample {question_id} has invalid response_type: {response_type!r}")
        contexts = sample.get("retrieved_contexts")
        if not isinstance(contexts, list) or not all(isinstance(context, str) for context in contexts):
            raise ValueError(f"Sample {question_id} has invalid retrieved contexts")
        if (response_type == "rag") != bool(contexts):
            raise ValueError(f"Sample {question_id} has inconsistent response_type and contexts")
        expected_trace_path = f"traces/{question_id}.json.gz"
        if sample.get("trace_path") != expected_trace_path:
            raise ValueError(f"Sample {question_id} has invalid trace_path: {sample.get('trace_path')!r}")
        validate_sample_trace(sample, traces[question_id], expected_trace_path)
    orphan_scores = sorted(set(scores) - set(samples))
    if orphan_scores:
        raise ValueError(f"Scores have no corresponding sample: {', '.join(orphan_scores)}")
    mismatched_response_types = sorted(
        question_id
        for question_id, score in scores.items()
        if score.get("response_type") != samples[question_id].get("response_type")
    )
    if mismatched_response_types:
        raise ValueError(f"Scores have mismatched response types: {', '.join(mismatched_response_types)}")

    ordered_traces = sorted(
        traces.items(),
        key=lambda item: (
            item[1].get("question_line") if isinstance(item[1].get("question_line"), int) else float("inf"),
            item[0],
        ),
    )
    return [
        _result_row(trace, samples.get(question_id), scores.get(question_id)) for question_id, trace in ordered_traces
    ]


def _style_header(cells) -> None:
    for cell in cells:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_results_sheet(workbook: Workbook, rows: list[Mapping[str, Any]]) -> None:
    worksheet = workbook.active
    worksheet.title = "Results"
    headers = [name for name, _description, _width in COLUMNS]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([_excel_value(row.get(header)) for header in headers])

    _style_header(worksheet[1])
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1

    for index, (_name, _description, width) in enumerate(COLUMNS, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    header_indexes = {cell.value: cell.column for cell in worksheet[1]}
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_number].height = 72
        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_cell = worksheet.cell(row_number, header_indexes["collection_status"])
        status_cell.fill = STATUS_FILLS.get(status_cell.value, PatternFill())
        classification_cell = worksheet.cell(row_number, header_indexes["classification"])
        classification_cell.fill = CLASSIFICATION_FILLS.get(classification_cell.value, PatternFill())
        for metric in METRIC_NAMES:
            worksheet.cell(row_number, header_indexes[metric]).number_format = "0.0000"

    table = Table(displayName="EvaluationResults", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    for metric in METRIC_NAMES:
        column = get_column_letter(header_indexes[metric])
        worksheet.conditional_formatting.add(
            f"{column}2:{column}{worksheet.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )


def _flatten_mapping(value: Mapping[str, Any], prefix: str = "") -> Iterable[tuple[str, Any]]:
    for key, item in value.items():
        path = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, Mapping):
            yield from _flatten_mapping(item, path)
        elif isinstance(item, list):
            yield path, json.dumps(item, ensure_ascii=False)
        else:
            yield path, item


def _append_summary_section(worksheet, title: str, value: Mapping[str, Any]) -> None:
    row_number = worksheet.max_row + 1
    worksheet.cell(row_number, 1, title)
    worksheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
    title_cell = worksheet.cell(row_number, 1)
    title_cell.fill = SECTION_FILL
    title_cell.font = Font(bold=True, color="1F1F1F")
    for field, field_value in _flatten_mapping(value):
        worksheet.append((_excel_value(field), _excel_value(field_value)))


def _write_summary_sheet(
    workbook: Workbook,
    run_dir: Path,
    manifest: Mapping[str, Any],
    summary: Mapping[str, Any],
) -> None:
    worksheet = workbook.create_sheet("Run Summary")
    worksheet.sheet_view.showGridLines = False
    worksheet.append(("source_run_directory", _excel_value(str(run_dir))))
    _append_summary_section(worksheet, "Run manifest", manifest)
    if summary:
        _append_summary_section(worksheet, "Evaluation summary", summary)
    worksheet.column_dimensions["A"].width = 42
    worksheet.column_dimensions["B"].width = 100
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_legend_sheet(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet("Legend")
    worksheet.sheet_view.showGridLines = False
    worksheet.append(("Column", "Description"))
    for name, description, _width in COLUMNS:
        worksheet.append((name, description))
    worksheet.append(())
    worksheet.append(("Route classification", "Internal graph value"))
    for raw, readable in CLASSIFICATION_LABELS.items():
        worksheet.append((readable, raw))
    _style_header(worksheet[1])
    classification_header_row = len(COLUMNS) + 3
    _style_header(worksheet[classification_header_row])
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 100
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_run_to_excel(
    run_dir: Path,
    output_path: Path | None = None,
    *,
    overwrite: bool = False,
) -> Path:
    run_dir = run_dir.expanduser().resolve()
    if not run_dir.is_dir():
        raise FileNotFoundError(f"Evaluation run directory not found: {run_dir}")

    manifest = _load_json_object(run_dir / "run.json", required=True)
    summary = _load_json_object(run_dir / "summary.json", required=False)
    rows = build_result_rows(run_dir)

    output_path = (output_path or run_dir / "evaluation_results.xlsx").expanduser().resolve()
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError(f"Excel output path must end in .xlsx: {output_path}")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Excel output already exists: {output_path}. Pass --overwrite to replace it.")
    output_path.parent.mkdir(mode=0o700, parents=True, exist_ok=True)

    workbook = Workbook()
    _write_results_sheet(workbook, rows)
    _write_summary_sheet(workbook, run_dir, manifest, summary)
    _write_legend_sheet(workbook)

    temporary_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        temporary_path.chmod(0o600)
        temporary_path.replace(output_path)
        output_path.chmod(0o600)
    finally:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
    return output_path
